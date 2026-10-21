"""Column selection for uploaded Quick Reports.

Each recognised report is one "tab" (General, Addresses, ...). `inspect` lists the columns of every tab so the user can
pick which ones the scan may use; `apply_selection` then drops the unticked columns from the loaded records.
A column is identified by its normalised header (spec_model.nk), the same key for supplier and customer reports.
"""
import csv
import io
import json

from openpyxl import load_workbook

from . import reports, spec_model
from .spec_model import nk

TAB_LABELS = {
    "general": "General", "address": "Addresses", "addrtype": "Address types", "comm": "Communication", "contact": "Contacts",
    "msg": "Messaging", "pay": "Payment", "purchase": "Purchasing", "invoice": "Invoicing", "srm": "SRM", "custord": "Customer orders",
}
TAB_ORDER = list(TAB_LABELS)
# Columns that tie a row to its supplier / customer / address; a scan cannot work without them.
LOCKED = {"supplierid", "customerid", "identity", "vendorno", "customerno", "addressid", "name", "party", "partytypedb"}


def _header_labels(filename: str, content: bytes) -> list[str]:
    if filename.lower().endswith(".csv"):
        first = next(csv.reader(io.StringIO(content.decode("utf-8-sig", errors="replace"))), [])
    else:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = next((w for w in wb.worksheets if w.max_row and w.max_row > 1), wb.worksheets[0])
        first = next(ws.iter_rows(values_only=True), ())
    return [str(h).strip() for h in first if h is not None and str(h).strip()]


def _tab(role: str, file: str, rows: int, labels: list[tuple[str, str]]) -> dict:
    seen, columns = set(), []
    for key, label in labels:
        if key and key not in seen and not reports.is_hidden(label):
            seen.add(key)
            columns.append({"key": key, "label": label, "locked": key in LOCKED})
    return {"role": role, "label": TAB_LABELS.get(role, role), "file": file, "rows": rows, "columns": columns}


def inspect(entity: str, uploads: list[tuple[str, bytes]]) -> list[dict]:
    """The recognised reports of the upload with their columns, in tab order."""
    return _inspect(entity, uploads)[0]


def _inspect(entity: str, uploads: list[tuple[str, bytes]]) -> tuple[list[dict], list[str]]:
    tabs: dict[str, dict] = {}
    notes: list[str] = []
    if entity == "supplier":
        tables, report = spec_model.load_tables(uploads)
        content = dict(uploads)
        for item in report:
            role = item["role"]
            if item["note"]:
                notes.append(f"{item['file']}: {item['note']}")
            if role not in tables or item["note"] or role in tabs:
                continue
            try:
                labels = [(nk(h), h) for h in _header_labels(item["file"], content[item["file"]])]
            except Exception:
                continue
            tabs[role] = _tab(role, item["file"], item["rows"], labels)
    else:
        res = reports.parse_uploads("customer", uploads)
        notes += [f"{u['file']}: {u['reason']}" for u in res["unmatched"]]
        role_of = {code: role for role, code in reports.ROLE_CODES["customer"].items()}
        stage_of = {c["code"]: c["stage"] for c in reports.catalog("customer")}
        for item in res["report"]:
            role = role_of.get(item["code"])
            records = res["stages"][stage_of[item["code"]]]["records"]
            if not role:
                continue
            labels = []
            for rec in records:
                labels += [(nk(k), k) for k in rec]
            tabs[role] = _tab(role, item["file"], item["rows"], labels)
    return [tabs[r] for r in TAB_ORDER if r in tabs], notes


def validate(entity: str, uploads: list[tuple[str, bytes]], selection: dict[str, set[str]] | None) -> dict:
    """Dry run of the upload and column choice: can the scan run, and what will it (not) use. Nothing is scored or stored."""
    tabs, notes = _inspect(entity, uploads)
    roles = {t["role"] for t in tabs}
    warnings = list(notes)
    if "general" not in roles:
        warnings.append(f"The general {entity} report is required (Supplier Id / Customer Id, Name, Association No ...).")
    expected = [r for r in TAB_ORDER if r in (reports.ROLE_CODES[entity])]
    missing = [TAB_LABELS[r] for r in expected if r not in roles]
    if missing:
        warnings.append("Reports not provided and skipped: " + ", ".join(missing) + ".")
    summary, ignored = [], 0
    for t in tabs:
        picked = {c["key"] for c in t["columns"] if c["locked"]} | (set(selection.get(t["role"], ())) if selection is not None else {c["key"] for c in t["columns"]})
        used = [c for c in t["columns"] if c["key"] in picked]
        ignored += len(t["columns"]) - len(used)
        summary.append({"role": t["role"], "label": t["label"], "rows": t["rows"], "columns": len(t["columns"]), "used": len(used)})
        if t["role"] != "general" and not any(not c["locked"] for c in used):
            warnings.append(f"{t['label']}: no columns selected, this report will not influence the scan.")
    general = next((t for t in tabs if t["role"] == "general"), None)
    if general and selection is not None and not any(not c["locked"] and c["key"] in selection.get("general", ()) for c in general["columns"]):
        warnings.append("General: only identifying columns are selected, so matching can rely on the name alone.")
    return {"valid": "general" in roles, "recordCount": general["rows"] if general else 0, "reports": summary,
            "ignoredColumns": ignored, "warnings": warnings}


def parse_selection(text: str) -> dict[str, set[str]] | None:
    """{role: {column key}} from the JSON the client sends; None when no selection was made (use everything)."""
    if not text:
        return None
    try:
        data = json.loads(text)
    except ValueError:
        raise ValueError("The column selection is not valid.")
    if not isinstance(data, dict):
        raise ValueError("The column selection is not valid.")
    return {str(role): {str(k) for k in keys} for role, keys in data.items() if isinstance(keys, list)}


def apply_selection(rows: list[dict], allowed: set[str] | None) -> list[dict]:
    """Records with only the selected columns (locked ones are always kept); rows left empty are dropped."""
    if allowed is None:
        return rows
    keep = allowed | LOCKED
    out = [{k: v for k, v in rec.items() if nk(k) in keep} for rec in rows]
    return [rec for rec in out if rec]
