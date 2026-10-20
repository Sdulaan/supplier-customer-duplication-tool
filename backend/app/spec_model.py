"""IFS supplier master deduplication, after "IFS Supplier Master Deduplication - Exact Table & Schema Specification".

Pipeline (section numbers refer to the spec):
  load the 10 Quick Report exports (3)  ->  flatten to one profile per Supplier Id (3)  ->  block candidate pairs
  ->  score each pair with the weighted matrix A01..P01 (5)  ->  tier it (6)  ->  connected components per lane
  ->  pick the Master Survivor of each cluster with the 100-point scorecard (6.1).

`spec_ml.py` adds a learned second opinion on top of the rule score.
"""
import csv
import io
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from itertools import combinations

from openpyxl import load_workbook

from .dedup import PUBLIC_DOMAINS, jaro_winkler, levenshtein, name_similarity, norm_address, phonetic

# --- spec constants -----------------------------------------------------------------------------------------------
W_A01, W_A02, W_N01, W_L01, W_C01, W_C02, W_B01 = 55, 45, 40, 30, 30, 15, 5
P_BRANCH = 25
TIER1, TIER2, TIER4 = 80, 60, 40
BOOSTER_GATE = 20
MAX_BLOCK = 300
ANCHOR_MAX_HOLDERS = 5  # an id shared by more suppliers than this is a placeholder, not an identity
COMM_MAX_HOLDERS = 6

LEGAL_SUFFIXES = {"ltd", "limited", "inc", "incorporated", "corp", "corporation", "pvt", "plc", "gmbh", "co", "company",
                  "llc", "sa", "ag", "bv", "nv", "pty", "llp", "lp", "se", "sdn", "bhd", "srl", "spa", "ab", "as", "oy"}
TAG_RE = re.compile(r"[-\s]*(?<!WITHOUT\s)\b(WHT|WITHHOLDINGS?(?:\s+TAX)?|SELF[\s-]*BILL(?:ING)?)\b", re.I)
JUNK_VALUES = ("TEST", "DUMMY", "NONE", "NULL", "UNKNOWN", "SAMPLE", "DEFAULT", "NOTAPPLICABLE", "TBC", "TBD", "XXXXX")
EDI_MEDIA = ("EDI", "E-INVOICE", "EINVOICE", "E_INVOICE")

TIER_NAMES = {1: "Auto-Confirm Duplicate", 2: "High-Probability Steward Review", 3: "Functional Split / Hierarchy Review",
              4: "Low Confidence / Audit Watchlist"}

# role -> columns that identify the report (normalized names), checked in order
ROLE_SIGNATURES = [
    ("purchase", {"vendorname"}), ("srm", {"extsourcesupplierid"}), ("invoice", {"invoicingsupplier"}),
    ("pay", {"blockedforpayment"}), ("addrtype", {"addresstypecodedb"}), ("msg", {"mediacode", "messageclass"}),
    ("comm", {"commid", "methodiddb"}), ("contact", {"personid", "guid"}), ("address", {"address1"}),
    ("general", {"associationno", "suppliercategorydb"}),
]
ROLE_FILE = {"general": "123054", "comm": "123124", "contact": "123145", "msg": "123211", "address": "123304",
             "addrtype": "123339", "pay": "123426", "purchase": "123442", "invoice": "123459", "srm": "123512"}


# --- loading ------------------------------------------------------------------------------------------------------
def nk(header) -> str:
    return re.sub(r"[^a-z0-9]", "", str(header).lower())


def _cell(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    s = str(v)
    return s if s.strip() else None  # keep leading/trailing spaces: ids that differ by a space are different records


def _rows(filename: str, content: bytes) -> tuple[set[str], list[dict]]:
    if filename.lower().endswith(".csv"):
        raw = list(csv.reader(io.StringIO(content.decode("utf-8-sig", errors="replace"))))
    else:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = next((w for w in wb.worksheets if w.max_row and w.max_row > 1), wb.worksheets[0])
        raw = list(ws.iter_rows(values_only=True))
    if not raw:
        return set(), []
    keys = [nk(h) if h is not None else "" for h in raw[0]]
    out = []
    for row in raw[1:]:
        rec = {k: c for k, v in zip(keys, row) if k and (c := _cell(v)) is not None}
        if rec:
            out.append(rec)
    return {k for k in keys if k}, out


def detect_role(columns: set[str]) -> str | None:
    for role, sig in ROLE_SIGNATURES:
        if sig <= columns:
            return role
    return None


def load_tables(files: list[tuple[str, bytes]]) -> tuple[dict[str, list[dict]], list[dict]]:
    """{role: records} from the uploaded exports, recognised by their columns; plus a report of what each file was taken as."""
    tables: dict[str, list[dict]] = {}
    report = []
    for name, content in files:
        try:
            headers, rows = _rows(name, content)
        except Exception as exc:
            report.append({"file": name, "role": None, "rows": 0, "note": f"unreadable ({exc})"})
            continue
        role = detect_role(headers)
        if not role:
            report.append({"file": name, "role": None, "rows": len(rows), "note": "not one of the 10 expected reports"})
            continue
        if role in tables and len(rows) <= len(tables[role]):
            report.append({"file": name, "role": role, "rows": len(rows), "note": "skipped, a larger file already fills this report"})
            continue
        tables[role] = rows
        report.append({"file": name, "role": role, "rows": len(rows), "note": ""})
    return tables, report


# --- value hygiene ------------------------------------------------------------------------------------------------
def alnum(v) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(v or "").upper())


def _sequential(s: str) -> bool:
    if len(s) < 4 or not (s.isdigit() or s.isalpha()):
        return False
    d = [ord(c) - ord(s[0]) for c in s]
    steps = {b - a for a, b in zip(d, d[1:])}
    return steps in ({1}, {-1}) or (s.isdigit() and steps <= {1, -9})


def valid_anchor(v) -> str:
    """Entropy filter (A01/A02): normalized value, or '' when it is a placeholder."""
    s = alnum(v)
    if len(s) < 5 or len(set(s)) < 3 or _sequential(s) or any(j in s for j in JUNK_VALUES):
        return ""
    return s


def norm_phone(v) -> str:
    d = re.sub(r"\D", "", str(v or ""))
    if len(d) < 7 or len(set(d)) < 3 or _sequential(d):
        return ""
    return d[-8:]


def norm_email(v) -> str:
    s = str(v or "").strip().lower()
    return s if re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", s) else ""


def clean_name(name: str) -> tuple[list[str], bool]:
    """(tokens without legal suffixes / functional tags / punctuation, had a WHT-style tag)."""
    s = str(name or "").upper().replace("&", " AND ")
    tagged = bool(TAG_RE.search(s))
    s = TAG_RE.sub(" ", s)
    toks = re.findall(r"[A-Z0-9]+", s)
    kept = [t for t in toks if t.lower() not in LEGAL_SUFFIXES and t != "AND"]
    return (kept or toks), tagged


def _initials(tokens: list[str]) -> str:
    return "".join(t[0] for t in tokens)


def acronym_match(ta: list[str], tb: list[str]) -> bool:
    for x, y in ((ta, tb), (tb, ta)):
        if len(x) == 1 and 2 <= len(x[0]) <= 6 and len(y) >= 2 and _initials(y) == x[0]:
            return True
        run = ""
        for t in x:
            if len(t) != 1:
                break
            run += t
        if len(run) >= 2 and len(y) >= len(run) and _initials(y[: len(run)]) == run:
            return True
    return False


# --- profiles -----------------------------------------------------------------------------------------------------
@dataclass
class Supplier:
    id: str
    name: str = ""
    created: str = ""
    country: str = ""
    category: str = ""            # SUPPLIER / PROSPECT
    b2b: bool = False
    one_time: bool = False
    template: bool = False
    internal: bool = False        # purchasing Category Db = 'I'
    quick: str = ""               # ORDINARY / QUICK
    buyer: str = ""
    supp_grp: str = ""
    currency: str = ""
    vendor_name: str = ""
    anchors: dict = field(default_factory=lambda: defaultdict(set))   # source label -> normalized values (A01)
    xrefs: dict = field(default_factory=lambda: defaultdict(set))     # source label -> normalized values (A02)
    phones: set = field(default_factory=set)
    emails: set = field(default_factory=set)
    domains: set = field(default_factory=set)
    contact_names: set = field(default_factory=set)
    persons: set = field(default_factory=set)
    addrs: list = field(default_factory=list)
    branches: set = field(default_factory=set)
    invoicing_suppliers: set = field(default_factory=set)
    other_payees: set = field(default_factory=set)
    rw_flags: set = field(default_factory=set)
    approval: set = field(default_factory=set)
    expired: bool = False
    in_invoice: bool = False
    in_pay: bool = False
    pay_blocked: bool = False
    in_purch: bool = False
    in_srm: bool = False
    # derived
    tokens: list = field(default_factory=list)
    clean: str = ""
    compact: str = ""
    phon: str = ""
    wht_tag: bool = False
    variants: list = field(default_factory=list)

    @property
    def id_clean(self) -> bool:
        return self.id == self.id.strip()


def _truthy(v) -> bool:
    return str(v or "").strip().upper() in ("TRUE", "YES", "Y", "1")


def build_suppliers(tables: dict[str, list[dict]]) -> dict[str, Supplier]:
    """Flatten the 10 reports to one profile per Supplier Id (spec section 3)."""
    sup: dict[str, Supplier] = {}
    for r in tables.get("general", []):
        sid = r.get("supplierid")
        if not sid:
            continue
        s = Supplier(sid)
        s.name = r.get("name", "")
        s.created = r.get("creationdate", "")
        s.country = (r.get("countrydb") or r.get("country") or "").upper()
        s.category = (r.get("suppliercategorydb") or "").upper()
        s.b2b = _truthy(r.get("b2bsupplierdb"))
        s.one_time = _truthy(r.get("onetimedb"))
        for col, label in (("associationno", "Association No"),):
            if v := valid_anchor(r.get(col)):
                s.anchors[label].add(v)
        if v := valid_anchor(r.get("suppliersownid")):
            s.xrefs["Suppliers Own Id"].add(v)
        sup[sid] = s

    for r in tables.get("purchase", []):
        s = sup.get(r.get("vendorno") or r.get("supplierid", ""))
        if not s:
            continue
        s.in_purch = True
        s.vendor_name = r.get("vendorname", "")
        s.template = (r.get("templatesupplierdb") or "").upper() == "TEMPLATE"
        s.internal = (r.get("categorydb") or "").upper() == "I"
        s.quick = (r.get("quickregisteredsupplierdb") or "").upper()
        s.buyer, s.supp_grp, s.currency = r.get("buyercode", ""), r.get("suppgrp", ""), (r.get("currencycode") or "").upper()
        for col in ("customerno", "ourcustomerno"):
            if v := valid_anchor(r.get(col)):
                s.xrefs["Customer No" if col == "customerno" else "Our Customer No"].add(v)

    def owner(r: dict) -> Supplier | None:
        if (r.get("partytypedb") or "SUPPLIER").upper() != "SUPPLIER":
            return None
        return sup.get(r.get("identity") or r.get("supplierid", ""))

    for r in tables.get("srm", []):
        if s := sup.get(r.get("supplierid", "")):
            s.in_srm = True
            if v := valid_anchor(r.get("extsourcesupplierid")):
                s.anchors["Ext Source Supplier Id"].add(v)

    addr_by_key: dict[tuple, dict] = {}
    for r in tables.get("address", []):
        s = sup.get(r.get("supplierid", ""))
        if not s:
            continue
        country = (r.get("countrydb") or r.get("country") or s.country).upper()
        na = norm_address(r.get("address1", ""), r.get("address2", ""), r.get("zipcode", ""), r.get("city", ""))
        a1 = " ".join(sorted(norm_address(r.get("address1", ""), "", "", "")["tokens"]))
        gln = valid_anchor(r.get("eanlocation"))
        branch = alnum(r.get("supplierbranch"))
        a = {"id": r.get("addressid", ""), "a1": a1 if len(a1) > 3 else "", "city": na["city"], "zip": na["zip"], "country": country,
             "gln": gln, "branch": branch, "text": na["text"], "city_text": na["city_text"], "types": set()}
        s.addrs.append(a)
        addr_by_key[(s.id, a["id"])] = a
        if branch:
            s.branches.add(branch)
        if not s.country:
            s.country = country
    for r in tables.get("addrtype", []):
        a = addr_by_key.get((r.get("supplierid", ""), r.get("addressid", "")))
        if a and _truthy(r.get("defaddress")):
            a["types"].add((r.get("addresstypecodedb") or "").upper())

    for r in tables.get("comm", []):
        s = owner(r)
        if not s:
            continue
        method, value = (r.get("methodiddb") or "").upper(), r.get("value", "")
        if method == "E_MAIL" or "@" in value:
            if e := norm_email(value):
                s.emails.add(e)
                s.domains.add(e.rsplit("@", 1)[1])
        elif method in ("PHONE", "MOBILE", ""):
            if p := norm_phone(value):
                s.phones.add(p)
        if n := alnum(r.get("name")):
            if len(n) >= 4:
                s.contact_names.add(n)

    for r in tables.get("contact", []):
        s = sup.get(r.get("supplierid", ""))
        if s and r.get("personid"):
            s.persons.add(r["personid"])

    for r in tables.get("msg", []):
        s = sup.get(r.get("supplierid", ""))
        if s and (r.get("mediacode") or "").upper() in EDI_MEDIA and (v := valid_anchor(r.get("address"))):
            s.xrefs["EDI / E-Invoice address"].add(v)

    for r in tables.get("pay", []):
        s = owner(r)
        if not s:
            continue
        s.in_pay = True
        s.pay_blocked = s.pay_blocked or _truthy(r.get("blockedforpayment"))
        if (o := r.get("otherpayeeidentity")) and o != s.id:
            s.other_payees.add(o)

    for r in tables.get("invoice", []):
        s = owner(r)
        if not s:
            continue
        s.in_invoice = True
        if (o := r.get("invoicingsupplier")) and o != s.id:
            s.invoicing_suppliers.add(o)
        for col, label in (("fiscalno", "Fiscal No"), ("supplierutr", "Supplier Utr"), ("secondtin", "Second Tin"),
                           ("certifiedemail", "Certified Email"), ("einvoicereceivercode", "Einvoice Receiver Code")):
            if v := valid_anchor(r.get(col)):
                s.anchors[label].add(v)
        if f := (r.get("reportandwithholddb") or "").upper():
            s.rw_flags.add(f)
        if st := (r.get("cfapprovalstatus") or "").upper():
            s.approval.add(st)
        if r.get("expiredate"):
            s.expired = True

    for s in sup.values():
        s.variants = []
        for nm in {s.name, s.vendor_name}:
            toks, _ = clean_name(nm)
            if toks:
                s.variants.append(toks)
        toks, s.wht_tag = clean_name(s.name or s.vendor_name)
        s.tokens = toks
        s.clean = " ".join(sorted(toks))
        s.compact = "".join(sorted(toks))
        s.phon = " ".join(sorted(c for c in (phonetic(t) for t in toks) if c))
    return sup


def eligible(s: Supplier) -> bool:
    """Template, one-time and internal (inter-site) suppliers never enter external duplicate matching (spec 2.1, 2.8, 3)."""
    return not (s.template or s.one_time or s.internal) and len(set(s.compact)) >= 3


# --- pair scoring -------------------------------------------------------------------------------------------------
@dataclass
class Context:
    """Dataset-wide holder counts, used to ignore identifiers and contacts shared by too many suppliers."""
    anchor_holders: Counter
    xref_holders: Counter
    phone_holders: Counter
    email_holders: Counter
    domain_holders: Counter
    name_holders: Counter
    person_holders: Counter
    gln_holders: Counter


def build_context(parties: list[Supplier]) -> Context:
    c = Context(*(Counter() for _ in range(8)))
    for s in parties:
        for v in set().union(*s.anchors.values()) if s.anchors else ():
            c.anchor_holders[v] += 1
        for v in set().union(*s.xrefs.values()) if s.xrefs else ():
            c.xref_holders[v] += 1
        c.phone_holders.update(s.phones)
        c.email_holders.update(s.emails)
        c.domain_holders.update(s.domains)
        c.name_holders.update(s.contact_names)
        c.person_holders.update(s.persons)
        c.gln_holders.update({a["gln"] for a in s.addrs if a["gln"]})
    return c


def name_score(a: Supplier, b: Supplier) -> tuple[int, float, str]:
    """N01: (points, best similarity, detail)."""
    best_sim, best_detail = 0.0, ""
    for ta in a.variants:
        for tb in b.variants:
            ca, cb = " ".join(sorted(ta)), " ".join(sorted(tb))
            if min(len("".join(ta)), len("".join(tb))) < 3:
                continue
            if ca == cb:
                return W_N01, 1.0, f'exact cleaned name "{ca}"'
            sim = max(name_similarity(ca, cb), jaro_winkler("".join(sorted(ta)), "".join(sorted(tb))))
            if sim > best_sim:
                best_sim, best_detail = sim, f'"{ca}" ~ "{cb}" ({sim:.2f})'
    if best_sim >= 0.95:
        return W_N01, best_sim, best_detail
    if best_sim >= 0.85:
        return 28, best_sim, best_detail
    if a.phon and b.phon and a.clean != b.clean and jaro_winkler(a.phon, b.phon) >= 0.88:
        return 18, best_sim, f"sounds alike: {a.phon} ~ {b.phon}"
    if acronym_match(a.tokens, b.tokens):
        return 18, best_sim, "acronym of the other name"
    return 0, best_sim, ""


def _street_sim(x: str, y: str) -> float:
    return jaro_winkler(x, y) if x and y else 0.0


def location_score(a: Supplier, b: Supplier, n01: int) -> tuple[int, float, str]:
    """L01: (points, best street similarity, detail), within the same Country Db."""
    best, street, why = 0, 0.0, ""
    gln_hit = None
    for x in a.addrs:
        for y in b.addrs:
            if x["country"] and y["country"] and x["country"] != y["country"]:
                continue
            same_city = bool(x["city"] and x["city"] == y["city"])
            same_zip = bool(x["zip"] and x["zip"] == y["zip"])
            sim = _street_sim(x["a1"], y["a1"])
            street = max(street, sim)
            pts = 0
            if x["a1"] and x["a1"] == y["a1"]:
                if same_city or same_zip:
                    pts, d = 30, f'{x["text"]}, {x["city_text"]}'
                elif not (x["city"] or x["zip"] or y["city"] or y["zip"]):
                    pts, d = 20, f'{x["text"]} (no city / zip on file)'  # same street line, locality unknown
                else:
                    pts, d = 0, ""
            elif sim >= 0.85 and (same_city or same_zip):
                pts, d = 20, f'{x["text"]} ~ {y["text"]}'
            elif same_city and same_zip and n01 >= 18:
                pts, d = 10, f'same city / zip {x["city_text"]} {x["zip"]}'.strip()
            if pts > best:
                best, why = pts, d
            if x["gln"] and x["gln"] == y["gln"]:
                gln_hit = x["gln"]
    if gln_hit and a.id != b.id:
        return min(30, best + 25), street, (why + "; " if why else "") + f"GLN {gln_hit}"
    return best, street, why


def contact_scores(a: Supplier, b: Supplier, ctx: Context, same_country: bool) -> tuple[int, int, list[str]]:
    """(C01 points, C02 points, details)."""
    c01, c02, notes = 0, 0, []
    ph = {p for p in a.phones & b.phones if ctx.phone_holders[p] <= COMM_MAX_HOLDERS}
    em = {e for e in a.emails & b.emails if ctx.email_holders[e] <= COMM_MAX_HOLDERS}
    if ph or em:
        c01 = W_C01
        notes.append("phone " + ", ".join(sorted(ph)) if ph else "email " + ", ".join(sorted(em)))
    else:
        dom = {d for d in a.domains & b.domains if d not in PUBLIC_DOMAINS and ctx.domain_holders[d] <= COMM_MAX_HOLDERS}
        nm = {n for n in a.contact_names & b.contact_names if ctx.name_holders[n] <= 3} if same_country else set()
        if dom or nm:
            c01 = 15
            notes.append("shared " + (f"domain {sorted(dom)[0]}" if dom else f"contact name {sorted(nm)[0]}"))
    for p in sorted(a.persons & b.persons):
        n = ctx.person_holders[p]
        pts = 15 if n <= 3 else 5 if n <= 6 else 0
        if pts > c02:
            c02 = pts
            notes.append(f"person {p} (linked to {n} suppliers)")
    return c01, c02, notes


def anchor_scores(a: Supplier, b: Supplier, ctx: Context) -> tuple[int, int, list[str], list[str]]:
    """(A01 points, A02 points, A01 details, A02 details)."""
    a01d, a02d = [], []
    for label, vals in a.anchors.items():
        for lb, vb in b.anchors.items():
            for v in vals & vb:
                if ctx.anchor_holders[v] <= ANCHOR_MAX_HOLDERS:
                    a01d.append(f"{label} = {v}" if label == lb else f"{label} = {lb} = {v}")
    for label, vals in a.xrefs.items():
        for lb, vb in b.xrefs.items():
            for v in vals & vb:
                if ctx.xref_holders[v] <= ANCHOR_MAX_HOLDERS:
                    a02d.append(f"{label} = {v}" if label == lb else f"{label} = {lb} = {v}")
    return (W_A01 if a01d else 0), (W_A02 if a02d else 0), sorted(set(a01d)), sorted(set(a02d))


def _conflicting_anchor(a: Supplier, b: Supplier) -> bool:
    """Both sides carry a registration number of the same kind and they clearly differ."""
    for label in set(a.anchors) & set(b.anchors):
        va, vb = a.anchors[label], b.anchors[label]
        if va and vb and not (va & vb) and all(levenshtein(x, y) > 2 for x in va for y in vb):
            return True
    return False


# Duplicate-checking conditions the user can switch on / off for a scan (name is always on, like Part Number / Description).
CONDITIONS = {"tax": ("A01", "Tax / registration number", "Association No, Fiscal No, Supplier Utr, Second Tin, Certified Email, E-invoice code, external registry id"),
              "xref": ("A02", "Cross-reference & EDI address", "Suppliers Own Id, Customer No, Our Customer No, EDI / E-Invoice address"),
              "address": ("L01", "Address & GLN", "Address1, City, Zip Code, Country, Ean Location"),
              "contact": ("C01", "Phone & email", "Communication method values, contact names, email domains"),
              "person": ("C02", "Linked contact person", "Person Id on the supplier contact table")}
ALL_CONDITIONS = tuple(CONDITIONS)


def score_pair(a: Supplier, b: Supplier, ctx: Context, conditions=ALL_CONDITIONS) -> dict:
    same_country = bool(a.country and a.country == b.country)
    a01, a02, a01d, a02d = anchor_scores(a, b, ctx)
    n01, name_sim, n01d = name_score(a, b)
    l01, street_sim, l01d = location_score(a, b, n01)
    c01, c02, cnotes = contact_scores(a, b, ctx, same_country)
    if "tax" not in conditions:
        a01, a01d = 0, []
    if "xref" not in conditions:
        a02, a02d = 0, []
    if "address" not in conditions:
        l01, street_sim, l01d = 0, 0.0, ""
    if "contact" not in conditions:
        c01 = 0
    if "person" not in conditions:
        c02 = 0
    if "contact" not in conditions or "person" not in conditions:
        cnotes = [n for n in cnotes if (("person" in n) == ("person" in conditions) or ("person" not in n and "contact" in conditions))]
    base = a01 + a02 + n01 + l01 + c01
    b01 = W_B01 if (base >= BOOSTER_GATE and same_country and a.currency and a.currency == b.currency
                    and ((a.buyer and a.buyer == b.buyer) or (a.supp_grp and a.supp_grp == b.supp_grp))) else 0

    tags = []
    if b.id in a.invoicing_suppliers | a.other_payees or a.id in b.invoicing_suppliers | b.other_payees:
        tags.append("LINKED HIERARCHY")
    if a.wht_tag != b.wht_tag or (a.rw_flags and b.rw_flags and a.rw_flags != b.rw_flags):
        tags.append("WHT FUNCTIONAL SPLIT")
    p_branch = -P_BRANCH if (a.branches and b.branches and not (a.branches & b.branches)) else 0

    raw = a01 + a02 + n01 + l01 + c01 + c02 + b01
    score = max(0, min(100, raw + p_branch))
    return {
        "score": score, "tags": tags,
        "rules": {"A01": a01, "A02": a02, "N01": n01, "L01": l01, "C01": c01, "C02": c02, "B01": b01, "P01": p_branch},
        "details": {"A01": a01d, "A02": a02d, "N01": n01d, "L01": l01d, "C": cnotes},
        "features": _features(a, b, a01, a02, n01, l01, c01, c02, b01, name_sim, street_sim, same_country, p_branch),
    }


FEATURE_NAMES = ["name_sim", "name_points", "phonetic_only", "street_sim", "loc_points", "c01", "c02", "same_country",
                 "same_currency", "same_buyer", "same_supp_grp", "branch_conflict", "wht_tag_differs", "both_have_address",
                 "id_sim"]


def _features(a: Supplier, b: Supplier, a01, a02, n01, l01, c01, c02, b01, name_sim, street_sim, same_country, p_branch) -> dict:
    """Feature vector for the ML layer. Anchors (A01/A02) are deliberately NOT features: they are what the model is
    trained to predict from the remaining evidence (see spec_ml)."""
    return {
        "name_sim": name_sim, "name_points": n01 / W_N01, "phonetic_only": float(n01 == 18), "street_sim": street_sim,
        "loc_points": l01 / 30, "c01": c01 / W_C01, "c02": c02 / W_C02, "same_country": float(same_country),
        "same_currency": float(bool(a.currency and a.currency == b.currency)),
        "same_buyer": float(bool(a.buyer and a.buyer == b.buyer)),
        "same_supp_grp": float(bool(a.supp_grp and a.supp_grp == b.supp_grp)),
        "branch_conflict": float(p_branch < 0), "wht_tag_differs": float(a.wht_tag != b.wht_tag),
        "both_have_address": float(bool(a.addrs and b.addrs)),
        "id_sim": jaro_winkler(alnum(a.id), alnum(b.id)),
    }


def tier_of(score: int, tags: list[str]) -> int | None:
    if tags:
        return 3 if score >= TIER4 else None
    if score >= TIER1:
        return 1
    if score >= TIER2:
        return 2
    if score >= TIER4:
        return 4
    return None


# --- candidate blocking -------------------------------------------------------------------------------------------
def candidate_pairs(parties: list[Supplier]) -> tuple[set[tuple[int, int]], int]:
    blocks: dict[tuple, list[int]] = defaultdict(list)
    index = {s.id: i for i, s in enumerate(parties)}
    direct: set[tuple[int, int]] = set()
    for i, s in enumerate(parties):
        for v in set().union(*s.anchors.values(), *s.xrefs.values()) if (s.anchors or s.xrefs) else ():
            blocks[("anchor", v)].append(i)
        if len(s.compact) >= 4:
            blocks[("name", s.compact[:5], s.country)].append(i)
        for t in s.tokens:
            if len(t) >= 4:
                blocks[("tok", t)].append(i)
        if s.phon:
            blocks[("ph", s.phon.split()[0])].append(i)
        for p in s.phones:
            blocks[("phone", p)].append(i)
        for e in s.emails:
            blocks[("mail", e)].append(i)
        for d in s.domains - PUBLIC_DOMAINS:
            blocks[("dom", d)].append(i)
        for p in s.persons:
            blocks[("person", p)].append(i)
        for a in s.addrs:
            if a["a1"]:
                blocks[("addr", a["a1"], a["country"], a["city"] or a["zip"])].append(i)
            if a["gln"]:
                blocks[("gln", a["gln"])].append(i)
        blocks[("idnorm", alnum(s.id))].append(i)
        for o in s.invoicing_suppliers | s.other_payees:
            if o in index and index[o] != i:
                direct.add(tuple(sorted((i, index[o]))))
    pairs, skipped = set(direct), 0
    for members in blocks.values():
        if len(members) < 2:
            continue
        if len(members) > MAX_BLOCK:
            skipped += 1
            continue
        pairs.update(combinations(sorted(set(members)), 2))
    return pairs, skipped


# --- survivor scorecard (6.1) -------------------------------------------------------------------------------------
def survivor_score(s: Supplier, earliest: str) -> tuple[int, dict]:
    parts = {"Party category": 25 if s.category == "SUPPLIER" else 0}
    inv = 0
    if s.in_invoice:
        if s.expired or "BLOCKED" in s.approval:
            inv = -30
        elif "ACTIVE" in s.approval:
            inv = 20
    parts["Invoice & approval"] = inv
    parts["Payment ledger active"] = 15 if (s.in_pay and not s.pay_blocked) else 0
    parts["Purchasing maturity"] = 15 if s.quick == "ORDINARY" else 5 if s.quick == "QUICK" else 0
    rich = (5 if s.anchors.get("Association No") else 0) + (5 if any(a["a1"] and a["city"] for a in s.addrs) else 0) \
        + (5 if (s.phones or s.emails) else 0)
    parts["Master data richness"] = rich
    parts["SRM present"] = 5 if s.in_srm else 0
    parts["Seniority / clean id"] = 5 if (s.created and s.created == earliest and s.id_clean) else 0
    return sum(parts.values()), parts


def _view(s: Supplier) -> dict:
    a = s.addrs[0] if s.addrs else {}
    return {"id": s.id, "name": s.name, "country": s.country, "category": s.category, "created": s.created,
            "address": a.get("text", ""), "city": a.get("city_text", ""), "phones": sorted(s.phones), "emails": sorted(s.emails),
            "anchors": {k: sorted(v) for k, v in s.anchors.items() if v}}


# --- orchestration ------------------------------------------------------------------------------------------------
def run(tables: dict[str, list[dict]], ml=None, conditions=ALL_CONDITIONS) -> dict:
    """Full pipeline. `ml` is an optional spec_ml.PairModel adding a learned probability to every scored pair."""
    sup = build_suppliers(tables)
    parties = [s for s in sup.values() if eligible(s)]
    ctx = build_context(parties)
    cand, skipped = candidate_pairs(parties)

    scored = []
    for i, j in cand:
        a, b = parties[i], parties[j]
        r = score_pair(a, b, ctx, conditions)
        r["ml"] = ml.predict(r["features"]) if ml else None
        r["tier"] = tier_of(r["score"], r["tags"])
        if r["tier"] is None and r["ml"] is not None and ml.can_surface and r["ml"] >= ml.surface_threshold and not r["tags"]:
            r["tier"], r["mlOnly"] = 4, True  # the rules see nothing, the model is confident
        if r["tier"]:
            scored.append((i, j, r))

    groups = _cluster(parties, scored)
    return {
        "groups": groups,
        "stats": {"suppliers": len(sup), "eligible": len(parties), "candidatePairs": len(cand), "skippedBlocks": skipped,
                  "scoredPairs": len(scored), "groups": len(groups), "conditions": list(conditions),
                  "byTier": dict(sorted(Counter(g["tier"] for g in groups).items())),
                  "excluded": {"template": sum(s.template for s in sup.values()), "oneTime": sum(s.one_time for s in sup.values()),
                               "internal": sum(s.internal for s in sup.values())}},
    }


def _cluster(parties: list[Supplier], scored: list) -> list[dict]:
    lanes = {"dup": [t for t in scored if t[2]["tier"] in (1, 2)], "hier": [t for t in scored if t[2]["tier"] == 3],
             "watch": [t for t in scored if t[2]["tier"] == 4]}
    groups = []
    for lane, pairs in lanes.items():
        parent: dict[int, int] = {}

        def find(x):
            parent.setdefault(x, x)
            while parent[x] != x:
                parent[x] = parent[parent[x]]
                x = parent[x]
            return x

        for i, j, _ in pairs:
            parent[find(i)] = find(j)
        members, comp_pairs = defaultdict(set), defaultdict(list)
        for i, j, r in pairs:
            root = find(i)
            members[root] |= {i, j}
            comp_pairs[root].append((i, j, r))
        for root, idx in members.items():
            ps = sorted(comp_pairs[root], key=lambda t: -t[2]["score"])
            ss = [parties[m] for m in idx]
            earliest = min((s.created for s in ss if s.created), default="")
            cards = {s.id: survivor_score(s, earliest) for s in ss}
            ranked = sorted(ss, key=lambda s: (-cards[s.id][0], s.created or "9999", s.id))
            tier = 1 if any(r["tier"] == 1 for _, _, r in ps) else ps[0][2]["tier"]
            top = ps[0][2]
            groups.append({
                "lane": lane, "tier": tier, "tierName": TIER_NAMES[tier], "score": top["score"], "tags": sorted({t for _, _, r in ps for t in r["tags"]}),
                "survivor": ranked[0].id if lane == "dup" else None,
                "records": [{**_view(s), "survivorScore": cards[s.id][0], "survivorBreakdown": cards[s.id][1]} for s in ranked],
                "pairs": [{"a": parties[i].id, "b": parties[j].id, "score": r["score"], "tier": r["tier"], "tags": r["tags"],
                           "rules": {k: v for k, v in r["rules"].items() if v}, "details": r["details"],
                           "ml": None if r["ml"] is None else round(r["ml"], 3), "mlOnly": r.get("mlOnly", False)}
                          for i, j, r in ps],
            })
    groups.sort(key=lambda g: ({"dup": 0, "hier": 1, "watch": 2}[g["lane"]], g["tier"], -g["score"], g["records"][0]["id"]))
    for n, g in enumerate(groups, 1):
        g["id"] = f"D{n:04d}"
    return groups


def run_files(files: list[tuple[str, bytes]], ml=None, conditions=ALL_CONDITIONS) -> dict:
    tables, report = load_tables(files)
    missing = [ROLE_FILE[r] + " " + r for r in ROLE_FILE if r not in tables]
    if "general" not in tables:
        raise ValueError("The SupplierInfoGeneral report (Supplier Id, Name, Association No, Supplier Category Db ...) is required.")
    out = run(tables, ml, conditions)
    out["files"] = report
    out["missingReports"] = missing
    return out
