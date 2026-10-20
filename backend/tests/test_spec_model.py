import io

from openpyxl import Workbook

from app import spec_ml, spec_model as sm


def xlsx(headers, rows):
    wb = Workbook()
    wb.active.append(headers)
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


GEN = ["Supplier Id", "Name", "Creation Date", "Association No", "Country Db", "Suppliers Own Id", "One Time Db", "Supplier Category Db"]
PUR = ["Vendor No", "Vendor Name", "Buyer Code", "Currency Code", "Supp Grp", "Category Db", "Template Supplier Db", "Quick Registered Supplier Db", "Customer No"]
ADR = ["Supplier Id", "Address Id", "Address1", "City", "Zip Code", "Country Db", "Supplier Branch", "Ean Location"]
COM = ["Identity", "Party Type Db", "Comm Id", "Method Id Db", "Value", "Name"]
INV = ["Company", "Identity", "Party Type Db", "Invoicing Supplier", "Report And Withhold Db", "Fiscal No", "Cf$ Approval Status", "Expire Date"]


def tables(gen, pur=(), adr=(), com=(), inv=()):
    files = [("g.xlsx", xlsx(GEN, gen))]
    if pur:
        files.append(("p.xlsx", xlsx(PUR, pur)))
    if adr:
        files.append(("a.xlsx", xlsx(ADR, adr)))
    if com:
        files.append(("c.xlsx", xlsx(COM, com)))
    if inv:
        files.append(("i.xlsx", xlsx(INV, inv)))
    return sm.load_tables(files)[0]


def pair(t):
    sup = sm.build_suppliers(t)
    ps = [s for s in sup.values() if sm.eligible(s)]
    return sm.score_pair(ps[0], ps[1], sm.build_context(ps))


def test_roles_detected_from_columns():
    _, rep = sm.load_tables([("x.xlsx", xlsx(GEN, [["1", "A", "2025-01-01", None, "US", None, "FALSE", "SUPPLIER"]])),
                             ("y.xlsx", xlsx(["Supplier Id", "Ext Source Supplier Id"], [["1", None]]))])
    assert {r["role"] for r in rep} == {"general", "srm"}


def test_exact_name_and_address_and_phone_is_tier1():
    t = tables([["S1", "Acme Trading Ltd", "2024-01-01", None, "US", None, "FALSE", "SUPPLIER"],
                ["S2", "ACME TRADING INC.", "2025-01-01", None, "US", None, "FALSE", "SUPPLIER"]],
               adr=[["S1", "1", "111 Main St", "Chicago", "60601", "US", None, None], ["S2", "1", "111 Main St", "Chicago", "60601", "US", None, None]],
               com=[["S1", "SUPPLIER", "1", "PHONE", "+1 312 555 0182", "x"], ["S2", "SUPPLIER", "1", "PHONE", "312-555-0182", "y"]])
    r = pair(t)
    assert r["score"] >= 80 and r["rules"]["N01"] == 40 and r["rules"]["L01"] == 30 and r["rules"]["C01"] == 30
    assert sm.tier_of(r["score"], r["tags"]) == 1


def test_dummy_association_numbers_are_not_anchors():
    t = tables([["S1", "Alpha Foods", "2024-01-01", "123", "US", None, "FALSE", "SUPPLIER"],
                ["S2", "Zenith Steel", "2024-01-01", "123", "US", None, "FALSE", "SUPPLIER"]])
    assert pair(t)["rules"]["A01"] == 0
    for v in ("12345", "11111", "TEST1", "NONE", "1234567890"):
        assert sm.valid_anchor(v) == ""
    assert sm.valid_anchor("IT123456") == "IT123456"


def test_valid_tax_id_is_anchor():
    t = tables([["S1", "Alpha Foods", "2024-01-01", "BRN-48213", "US", None, "FALSE", "SUPPLIER"],
                ["S2", "Alpha Food", "2024-01-01", "brn 48213", "US", None, "FALSE", "SUPPLIER"]])
    r = pair(t)
    assert r["rules"]["A01"] == 55 and r["score"] >= 80


def test_wht_split_is_tier3_not_auto_merge():
    t = tables([["GL10", "GL Supplier", "2024-01-01", None, "GB", None, "FALSE", "SUPPLIER"],
                ["GL11-WHT", "GL Supplier - WHT", "2024-01-01", None, "GB", None, "FALSE", "SUPPLIER"]],
               com=[["GL10", "SUPPLIER", "1", "PHONE", "+44 26389474", "x"], ["GL11-WHT", "SUPPLIER", "1", "PHONE", "+4426389474", "y"]])
    r = pair(t)
    assert r["tags"] == ["WHT FUNCTIONAL SPLIT"] and r["score"] >= 70
    assert sm.tier_of(r["score"], r["tags"]) == 3


def test_without_withholding_is_not_a_wht_tag():
    toks, tagged = sm.clean_name("Annick Account Payable without withholding")
    assert not tagged and "WITHOUT" in toks


def test_linked_hierarchy_tag():
    t = tables([["56", "Beta Corp", "2024-01-01", None, "US", None, "FALSE", "SUPPLIER"],
                ["SUP001", "Beta Corporation", "2024-01-01", None, "US", None, "FALSE", "SUPPLIER"]],
               inv=[["C1", "56", "SUPPLIER", "SUP001", "BLOCKED", None, "Active", None]])
    assert "LINKED HIERARCHY" in pair(t)["tags"]


def test_distinct_branches_are_penalised():
    t = tables([["S1", "Singer Lanka", "2024-01-01", None, "IN", None, "FALSE", "SUPPLIER"],
                ["S2", "Singer Lanka", "2024-01-01", None, "IN", None, "FALSE", "SUPPLIER"]],
               adr=[["S1", "1", "5 Road", "Mumbai", "400001", "IN", "MUMBAI", None], ["S2", "1", "5 Road", "Mumbai", "400001", "IN", "LONDON", None]])
    assert pair(t)["rules"]["P01"] == -25


def test_excluded_suppliers_never_match():
    t = tables([["S1", "Same Name", "2024-01-01", None, "US", None, "FALSE", "SUPPLIER"], ["S2", "Same Name", "2024-01-01", None, "US", None, "FALSE", "SUPPLIER"]],
               pur=[["S1", "Same Name", "B", "USD", "10", "E", "TEMPLATE", "ORDINARY", None], ["S2", "Same Name", "B", "USD", "10", "E", "NOTEMPLATE", "ORDINARY", None]])
    assert len(sm.run(t)["groups"]) == 0


def test_survivor_prefers_supplier_active_ordinary():
    t = tables([["S1", "Acme Trading", "2025-06-01", "BRN-99887", "US", None, "FALSE", "PROSPECT"],
                ["S2", "Acme Trading", "2024-01-01", "BRN-99887", "US", None, "FALSE", "SUPPLIER"]],
               pur=[["S1", "Acme Trading", "B", "USD", "10", "E", "NOTEMPLATE", "QUICK", None], ["S2", "Acme Trading", "B", "USD", "10", "E", "NOTEMPLATE", "ORDINARY", None]],
               inv=[["C1", "S2", "SUPPLIER", "S2", "BLOCKED", None, "Active", None]])
    g = sm.run(t)["groups"][0]
    assert g["tier"] == 1 and g["survivor"] == "S2"


def test_shared_generic_person_is_ignored():
    ctx = sm.Context(*(sm.Counter() for _ in range(8)))
    ctx.person_holders["147"] = 10
    a, b = sm.Supplier("A"), sm.Supplier("B")
    a.persons = b.persons = {"147"}
    assert sm.contact_scores(a, b, ctx, True)[1] == 0
    ctx.person_holders["147"] = 2
    assert sm.contact_scores(a, b, ctx, True)[1] == 15


def test_model_trains_and_scores():
    gen, adr = [], []
    for i in range(30):  # 30 anchor-sharing pairs (positives) and 30 conflicting-anchor pairs (negatives)
        for k, name, brn in ((0, f"Firm{i}x Trading", f"REG{i:05d}1"), (1, f"Firm{i}x Trade", f"REG{i:05d}1")):
            gen.append([f"P{i}-{k}", name, "2024-01-01", brn, "US", None, "FALSE", "SUPPLIER"])
            adr.append([f"P{i}-{k}", "1", f"{i} Main St", "Town", "1000", "US", None, None])
        for k, brn in ((0, f"AAA{i:05d}7"), (1, f"ZZZ{i:05d}3")):
            gen.append([f"N{i}-{k}", f"Firm{i}y Trading", "2024-01-01", brn, "US", None, "FALSE", "SUPPLIER"])
    t = tables(gen, adr=adr)
    model = spec_ml.train(t)
    assert model.meta["positives"] >= 30
    out = sm.run(t, model)
    assert out["groups"] and all(g["pairs"][0]["ml"] is not None for g in out["groups"])


def test_scan_store_roundtrip(tmp_path, monkeypatch):
    from app import spec_export, spec_store
    monkeypatch.setattr(spec_store, "DB_PATH", tmp_path / "scans.db")
    t = tables([["S1", "Acme Trading", "2024-01-01", "BRN-99887", "US", None, "FALSE", "SUPPLIER"],
                ["S2", "Acme Trading", "2024-01-01", "BRN-99887", "US", None, "FALSE", "SUPPLIER"]])
    out = sm.run(t)
    out["files"] = []
    saved = spec_store.save_scan(out, "First")
    assert [s["name"] for s in spec_store.list_scans()] == ["First"]
    row, result = spec_store.get_scan(saved["id"])
    assert result["stats"]["groups"] == 1 and row["byTier"]
    assert spec_export.build_report(result)[:2] == b"PK"
    assert spec_store.delete_scan(saved["id"]) and spec_store.list_scans() == []


def test_decisions_saved_and_in_report(tmp_path, monkeypatch):
    import io
    from openpyxl import load_workbook
    from app import spec_export, spec_store
    monkeypatch.setattr(spec_store, "DB_PATH", tmp_path / "scans.db")
    t = tables([["S1", "Acme Trading", "2024-01-01", "BRN-99887", "US", None, "FALSE", "SUPPLIER"],
                ["S2", "Acme Trading", "2024-01-01", "BRN-99887", "US", None, "FALSE", "SUPPLIER"]])
    out = sm.run(t)
    out["files"] = []
    sid = spec_store.save_scan(out)["id"]
    gid = out["groups"][0]["id"]
    assert spec_store.set_decision(sid, gid, "Confirmed as duplicate", "clean up")["decision"] == "Confirmed as duplicate"
    assert spec_store.set_decision(99, gid, "Deferred") is None
    try:
        spec_store.set_decision(sid, gid, "Maybe")
        assert False
    except ValueError:
        pass
    wb = load_workbook(io.BytesIO(spec_export.build_report(out, decisions=spec_store.get_decisions(sid))))
    ws = wb["Group Details"]
    assert ws["J2"].value == "Confirmed as duplicate" and ws["K2"].value == "clean up"
    spec_store.set_decision(sid, gid, "")
    assert spec_store.get_decisions(sid) == {}


def test_conditions_switch_rules_off():
    t = tables([["S1", "Acme Trading", "2024-01-01", "BRN-99887", "US", None, "FALSE", "SUPPLIER"],
                ["S2", "Acme Trading", "2024-01-01", "BRN-99887", "US", None, "FALSE", "SUPPLIER"]])
    sup = [s for s in sm.build_suppliers(t).values()]
    ctx = sm.build_context(sup)
    assert sm.score_pair(sup[0], sup[1], ctx)["rules"]["A01"] == 55
    r = sm.score_pair(sup[0], sup[1], ctx, conditions=("address",))
    assert r["rules"]["A01"] == 0 and r["rules"]["N01"] == 40
    assert sm.run(t, conditions=())["stats"]["conditions"] == []


def test_fetched_supplier_maps_to_tables_and_scans():
    from app import scan_adapters
    fetched = {"entity": "supplier", "stages": {
        "general": {"records": [{"SupplierId": "S1", "Name": "Acme Trading", "AssociationNo": "BRN-7788", "CountryDb": "US"},
                                {"SupplierId": "S2", "Name": "Acme Trading Inc", "AssociationNo": "BRN-7788", "CountryDb": "US"}]},
        "SupplierInfoAddresses": {"records": [{"_SupplierId": "S1", "AddressId": "1", "Address1": "1 Main St", "City": "X"}]},
        "SupplierInfoAddresses/AddressCommunicationMethods": {"records": [{"_SupplierId": "S1", "MethodIdDb": "E_MAIL", "Value": "a@acme.com"}]}}}
    t = scan_adapters.tables_from_fetched(fetched)
    assert {k: len(v) for k, v in t.items()} == {"general": 2, "address": 1, "comm": 1}
    assert sm.run(t)["stats"]["byTier"] == {1: 1}


def test_customer_scan_reshapes_groups():
    from app import scan_adapters
    key = {"entity": "customer", "stages": {"general": {"records": [
        {"CustomerId": "C1", "Name": "Acme Trading Ltd", "AssociationNo": "BRN-12345", "Country": "US"},
        {"CustomerId": "C2", "Name": "Acme Trading Limited", "AssociationNo": "BRN-12345", "Country": "US"}]}}}
    out = scan_adapters.customer_scan([key])
    g = out["groups"][0]
    assert out["entity"] == "customer" and g["tier"] in (1, 2) and g["survivor"] == "C1" and {r["id"] for r in g["records"]} == {"C1", "C2"}
